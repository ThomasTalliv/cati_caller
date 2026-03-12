"""Excel export using openpyxl."""
from __future__ import annotations

import os
import uuid
from datetime import datetime, timezone
from typing import Any

from cati.export.base import Exporter, ExportResult


class ExcelExporter(Exporter):
    """Export survey data to a styled Excel workbook."""

    def __init__(self, export_dir: str) -> None:
        self._export_dir = export_dir

    async def export(
        self,
        survey_id: uuid.UUID,
        *,
        call_ids: list[uuid.UUID] | None = None,
    ) -> ExportResult:
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError as exc:
            raise RuntimeError("openpyxl not installed. Run: pip install openpyxl") from exc

        from cati.db.session import get_session_factory
        from cati.db.repositories.response_repo import ResponseRepository
        from cati.db.repositories.call_repo import CallRepository
        from cati.db.repositories.survey_repo import SurveyRepository

        factory = get_session_factory()

        async with factory() as session:
            survey_repo = SurveyRepository(session)
            call_repo = CallRepository(session)
            survey = await survey_repo.get(survey_id)
            calls = await call_repo.list(survey_id=survey_id, status="completed", limit=10000)
            if call_ids:
                calls = [c for c in calls if c.id in set(call_ids)]

        question_keys = [q.question_key for q in survey.questions] if survey else []

        wb = openpyxl.Workbook()

        # ── Sheet 1: Summary ──────────────────────────────────────────────────
        ws_summary = wb.active
        ws_summary.title = "Summary"
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(fill_type="solid", fgColor="2E75B6")

        ws_summary.append(["Survey Export"])
        ws_summary.append(["Survey ID", str(survey_id)])
        ws_summary.append(["Generated", datetime.now(timezone.utc).isoformat()])
        ws_summary.append(["Total Completed Calls", len(calls)])
        ws_summary.append([])

        # ── Sheet 2: Responses ────────────────────────────────────────────────
        ws_responses = wb.create_sheet("Responses")
        headers = ["Call ID", "Phone", "Duration (s)", "Date"] + question_keys
        ws_responses.append(headers)
        for cell in ws_responses[1]:
            cell.font = header_font
            cell.fill = header_fill

        async with factory() as session:
            resp_repo = ResponseRepository(session)
            for call in calls:
                responses = await resp_repo.list_for_call(call.id)
                resp_map = {r.question_key: r.parsed_value for r in responses}
                row = [
                    str(call.id),
                    call.phone_number,
                    call.duration_s or 0,
                    call.ended_at.isoformat() if call.ended_at else "",
                ] + [resp_map.get(k, "") for k in question_keys]
                ws_responses.append(row)

        # ── Sheet 3: Transcripts ──────────────────────────────────────────────
        ws_transcripts = wb.create_sheet("Transcripts")
        ws_transcripts.append(["Call ID", "Speaker", "Text", "Timestamp"])
        for cell in ws_transcripts[1]:
            cell.font = header_font
            cell.fill = header_fill

        async with factory() as session:
            resp_repo = ResponseRepository(session)
            for call in calls:
                turns = await resp_repo.list_transcript(call.id)
                for t in turns:
                    ws_transcripts.append([str(call.id), t.speaker, t.text, t.created_at.isoformat()])

        os.makedirs(self._export_dir, exist_ok=True)
        filename = f"survey_{survey_id}_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}.xlsx"
        path = os.path.join(self._export_dir, filename)
        wb.save(path)

        return ExportResult(
            export_id=uuid.uuid4(),
            format="excel",
            file_path=path,
            call_count=len(calls),
        )

    @property
    def format_name(self) -> str:
        return "excel"
