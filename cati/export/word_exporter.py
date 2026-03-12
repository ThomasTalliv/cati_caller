"""Word document export using python-docx."""
from __future__ import annotations

import os
import uuid
from datetime import datetime, timezone

from cati.export.base import Exporter, ExportResult


class WordExporter(Exporter):
    """Export survey data to a formatted Word document."""

    def __init__(self, export_dir: str) -> None:
        self._export_dir = export_dir

    async def export(
        self,
        survey_id: uuid.UUID,
        *,
        call_ids: list[uuid.UUID] | None = None,
    ) -> ExportResult:
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError as exc:
            raise RuntimeError("python-docx not installed. Run: pip install python-docx") from exc

        from cati.db.session import get_session_factory
        from cati.db.repositories.response_repo import ResponseRepository
        from cati.db.repositories.call_repo import CallRepository
        from cati.db.repositories.survey_repo import SurveyRepository

        factory = get_session_factory()
        async with factory() as session:
            survey_repo = SurveyRepository(session)
            call_repo = CallRepository(session)
            survey = await survey_repo.get(survey_id)
            calls = await call_repo.list(survey_id=survey_id, status="completed", limit=10000)
            if call_ids:
                calls = [c for c in calls if c.id in set(call_ids)]

        doc = Document()

        # Title
        title = doc.add_heading("CATI Survey Report", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Metadata table
        meta = doc.add_table(rows=3, cols=2)
        meta.style = "Table Grid"
        meta.cell(0, 0).text = "Survey ID"
        meta.cell(0, 1).text = str(survey_id)
        meta.cell(1, 0).text = "Survey Name"
        meta.cell(1, 1).text = survey.name if survey else "Unknown"
        meta.cell(2, 0).text = "Generated"
        meta.cell(2, 1).text = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
        doc.add_paragraph()

        # Summary
        doc.add_heading("Summary", level=1)
        doc.add_paragraph(f"Total completed calls: {len(calls)}")
        doc.add_paragraph()

        # Per-call sections
        async with factory() as session:
            resp_repo = ResponseRepository(session)
            for i, call in enumerate(calls, start=1):
                doc.add_heading(f"Call {i}: {call.phone_number}", level=2)
                doc.add_paragraph(f"Call ID: {call.id}")
                doc.add_paragraph(f"Date: {call.ended_at.isoformat() if call.ended_at else 'N/A'}")
                doc.add_paragraph(f"Duration: {call.duration_s or 0} seconds")

                responses = await resp_repo.list_for_call(call.id)
                if responses:
                    doc.add_heading("Responses", level=3)
                    tbl = doc.add_table(rows=1, cols=3)
                    tbl.style = "Table Grid"
                    hdr = tbl.rows[0].cells
                    hdr[0].text = "Question Key"
                    hdr[1].text = "Answer"
                    hdr[2].text = "Refused"
                    for r in responses:
                        row = tbl.add_row().cells
                        row[0].text = r.question_key
                        row[1].text = str(r.parsed_value or "")
                        row[2].text = "Yes" if r.is_refused else "No"

                turns = await resp_repo.list_transcript(call.id)
                if turns:
                    doc.add_heading("Transcript", level=3)
                    for t in turns:
                        p = doc.add_paragraph()
                        run = p.add_run(f"[{t.speaker.upper()}] ")
                        run.bold = True
                        p.add_run(t.text)

                doc.add_page_break()

        os.makedirs(self._export_dir, exist_ok=True)
        filename = f"survey_{survey_id}_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}.docx"
        path = os.path.join(self._export_dir, filename)
        doc.save(path)

        return ExportResult(
            export_id=uuid.uuid4(),
            format="word",
            file_path=path,
            call_count=len(calls),
        )

    @property
    def format_name(self) -> str:
        return "word"
